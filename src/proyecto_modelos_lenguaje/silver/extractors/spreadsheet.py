"""Spreadsheet extraction with openpyxl."""

from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook

from ..models import ExtractedBlock, ExtractedDocument


def extract_xlsx(data: bytes, *, title: str | None = None) -> ExtractedDocument:
    workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    blocks: list[ExtractedBlock] = []
    sheets: list[str] = []
    try:
        for sheet in workbook.worksheets:
            rows: list[str] = []
            for row in sheet.iter_rows(values_only=True):
                cells = [
                    str(value).strip() for value in row if value is not None and str(value).strip()
                ]
                if cells:
                    rows.append(" | ".join(cells))
            if not rows:
                continue
            sheets.append(sheet.title)
            blocks.append(ExtractedBlock(text=sheet.title, level=2))
            blocks.append(ExtractedBlock(text="\n".join(rows)))
    finally:
        workbook.close()
    if title:
        blocks.insert(0, ExtractedBlock(text=title, level=1))
    return ExtractedDocument(
        title=title,
        blocks=blocks,
        metadata={"sheets": sheets},
        skip_reason=None if sheets else "no_text",
    )
