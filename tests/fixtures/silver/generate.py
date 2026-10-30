"""Generate binary fixtures at test time; no binary files are versioned."""

from __future__ import annotations

import contextlib
from pathlib import Path


def generate_pdf(
    path: Path,
    *,
    pages: int = 2,
    text: bool = True,
    title: str = "Informe de prueba",
) -> Path:
    from fpdf import FPDF

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    for number in range(1, pages + 1):
        pdf.add_page()
        if not text:
            continue
        if number == 1:
            with contextlib.suppress(Exception):
                pdf.start_section(title, level=1)
        pdf.set_font("Helvetica", "B", 15)
        pdf.cell(0, 10, f"Seccion {number}", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", 10)
        pdf.cell(0, 5, "Codigo: FR-PP-001", new_x="LMARGIN", new_y="NEXT")
        pdf.cell(0, 5, f"Pagina {number} de {pages}", new_x="LMARGIN", new_y="NEXT")
        if text:
            pdf.multi_cell(
                0,
                6,
                "Manuelita es una organizacion agroindustrial diversificada. "
                "La compania opera en Colombia, Peru y Chile y promueve el desarrollo "
                "sostenible de las comunidades vecinas. " * 3,
            )
    path.parent.mkdir(parents=True, exist_ok=True)
    pdf.output(str(path))
    return path


def generate_encrypted_pdf(path: Path) -> Path:
    from fpdf import FPDF

    pdf = FPDF()
    pdf.set_encryption(user_password="secret", owner_password="owner")
    pdf.add_page()
    pdf.set_font("Helvetica", "", 11)
    pdf.multi_cell(0, 6, "Contenido protegido de Manuelita.")
    path.parent.mkdir(parents=True, exist_ok=True)
    pdf.output(str(path))
    return path


def generate_xlsx(path: Path) -> Path:
    from openpyxl import Workbook

    workbook = Workbook()
    products = workbook.active
    products.title = "Productos"
    products.append(["Producto", "Unidad", "Descripcion"])
    products.append(["Azucar", "kg", "Azucar refinada de cana"])
    products.append(["Bioetanol", "litro", "Combustible renovable"])
    sites = workbook.create_sheet("Sedes")
    sites.append(["Sede", "Pais"])
    sites.append(["Valle del Cauca", "Colombia"])
    sites.append(["Trujillo", "Peru"])
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path
